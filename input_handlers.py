
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path

import pymupdf
from PIL import Image
from docx import Document as WordDocument
from docx.text.paragraph import Paragraph
from docx.table import Table
from openpyxl import load_workbook
from datetime import date, datetime, time

@dataclass
class DocumentPart:
    label: str
    text: str = ""
    image_bytes: bytes | None = None
    structured_data: dict | None = None


IMAGE_EXTENSIONS = {
    ".png", ".jpg", ".jpeg", ".webp",
    ".bmp", ".tif", ".tiff"
}


def _image_to_png(path: Path) -> bytes:
    """Convert an image into PNG bytes for consistent model input."""
    with Image.open(path) as image:
        rgb_image = image.convert("RGB")
        buffer = BytesIO()
        rgb_image.save(buffer, format="PNG")
        return buffer.getvalue()


def _load_pdf(path: Path) -> list[DocumentPart]:
    parts = []

    with pymupdf.open(path) as pdf:
        if len(pdf) == 0:
            raise ValueError("The PDF contains no pages.")

        for page_number, page in enumerate(pdf, start=1):
            # Keep extracted text when available.
            text = page.get_text("text").strip()

            # Render the page so Gemma can inspect its visual layout.
            pixmap = page.get_pixmap(
                matrix=pymupdf.Matrix(2, 2),
                alpha=False,
            )
            image_bytes = pixmap.tobytes("png")

            parts.append(
                DocumentPart(
                    label=f"Page {page_number}",
                    text=text,
                    image_bytes=image_bytes,
                )
            )

    return parts


def _load_image(path: Path) -> list[DocumentPart]:
    return [
        DocumentPart(
            label="Image 1",
            image_bytes=_image_to_png(path),
        )
    ]


def _load_docx(path: Path) -> list[DocumentPart]:
    document = WordDocument(path)
    lines = []
    table_number = 0

    # Preserve the order of paragraphs and tables.
    for block in document.iter_inner_content():
        if isinstance(block, Paragraph):
            text = block.text.strip()
            if text:
                lines.append(text)

        elif isinstance(block, Table):
            table_number += 1
            lines.append(f"[Table {table_number}]")

            for row in block.rows:
                cells = [
                    cell.text.replace("\n", " / ").strip()
                    for cell in row.cells
                ]

                if any(cells):
                    lines.append(" | ".join(cells))

    text = "\n".join(lines)

    if not text.strip():
        text = "[No paragraph or table text extracted.]"

    return [DocumentPart(label="Word document", text=text)]


def _load_txt(path: Path) -> list[DocumentPart]:
    try:
        text = path.read_text(encoding="utf-8-sig")
    except UnicodeDecodeError:
        text = path.read_text(encoding="cp1252")

    return [DocumentPart(label="Text document", text=text)]



def _load_xlsx(path: Path) -> list[DocumentPart]:
    """Read spreadsheet cells directly and preserve their structure."""

    workbook = load_workbook(
        path,
        read_only=True,
        data_only=True,
    )

    parts = []

    def safe_value(value):
        if value is None:
            return None

        if isinstance(value, (datetime, date, time)):
            return value.isoformat()

        if isinstance(value, (str, int, float, bool)):
            return value

        return str(value)

    try:
        for sheet in workbook.worksheets:
            iterator = iter(sheet.iter_rows(values_only=True))

            # Find the first non-empty row and treat it as the header.
            header_row = None
            header_row_number = None

            for row_number, row in enumerate(iterator, start=1):
                if any(
                    value is not None and str(value).strip()
                    for value in row
                ):
                    header_row = row
                    header_row_number = row_number
                    break

            if header_row is None:
                parts.append(
                    DocumentPart(
                        label=f"Sheet '{sheet.title}'",
                        structured_data={
                            "title": sheet.title,
                            "headers": [],
                            "rows": [],
                        },
                    )
                )
                continue

            headers = []

            for column_number, value in enumerate(header_row, start=1):
                header = safe_value(value)

                if header is None or not str(header).strip():
                    header = f"Column {column_number}"

                headers.append(str(header))

            # Read the remaining rows, preserving original cell values.
            data_rows = []

            for row_number, row in enumerate(
                iterator,
                start=header_row_number + 1,
            ):
                values = [
                    safe_value(value)
                    for value in row[:len(headers)]
                ]

                values.extend(
                    [None] * (len(headers) - len(values))
                )

                if not any(
                    value is not None and str(value).strip()
                    for value in values
                ):
                    continue

                data_rows.append({
                    "row_number": row_number,
                    "values": values,
                })

            # Keep chunks manageable for later AI analysis.
            chunk_size = 25

            if not data_rows:
                data_rows = []

            chunks = [
                data_rows[i:i + chunk_size]
                for i in range(0, len(data_rows), chunk_size)
            ] or [[]]

            for chunk_number, chunk in enumerate(chunks, start=1):
                row_numbers = [
                    item["row_number"] for item in chunk
                ]

                row_start = row_numbers[0] if row_numbers else header_row_number
                row_end = row_numbers[-1] if row_numbers else header_row_number

                parts.append(
                    DocumentPart(
                        label=(
                            f"Sheet '{sheet.title}', "
                            f"rows {row_start}-{row_end}"
                        ),
                        structured_data={
                            "title": sheet.title,
                            "headers": headers,
                            "rows": [
                                item["values"] for item in chunk
                            ],
                            "source_row_numbers": row_numbers,
                            "chunk": chunk_number,
                        },
                    )
                )

    finally:
        workbook.close()

    return parts


    parts = []

    try:
        for sheet in workbook.worksheets:
            # Process up to 50 worksheet rows per part.
            for start_row in range(1, sheet.max_row + 1, 50):
                end_row = min(start_row + 49, sheet.max_row)
                lines = []

                for row in sheet.iter_rows(
                    min_row=start_row,
                    max_row=end_row,
                ):
                    cells = []

                    for cell in row:
                        if cell.value is not None:
                            value = str(cell.value).replace("\n", " / ")
                            cells.append(f"{cell.coordinate}={value}")

                    if cells:
                        lines.append(" | ".join(cells))

                if lines:
                    parts.append(
                        DocumentPart(
                            label=(
                                f"Sheet '{sheet.title}', "
                                f"rows {start_row}-{end_row}"
                            ),
                            text="\n".join(lines),
                        )
                    )
    finally:
        workbook.close()

    if not parts:
        parts.append(
            DocumentPart(
                label="Workbook",
                text="[No populated cells found.]",
            )
        )

    return parts


def load_document(path: str | Path) -> tuple[str, list[DocumentPart]]:
    """Return the detected file type and its prepared document parts."""
    path = Path(path).expanduser()

    if not path.is_file():
        raise FileNotFoundError(f"File not found: {path}")

    extension = path.suffix.lower()

    if extension == ".pdf":
        return "pdf", _load_pdf(path)

    if extension in IMAGE_EXTENSIONS:
        return "image", _load_image(path)

    if extension == ".docx":
        return "docx", _load_docx(path)

    if extension == ".txt":
        return "txt", _load_txt(path)

    if extension == ".xlsx":
        return "xlsx", _load_xlsx(path)

    raise ValueError(
        f"Unsupported file type: {extension}\n"
        "Supported formats: PDF, common images, DOCX, TXT, XLSX."
    )
